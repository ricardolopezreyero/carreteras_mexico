# -*- coding: utf-8 -*-
# RLR · Perfiles de puesto y manuales de ingreso (Word) — Ricardo López Reyero · EYE 181218
# Un documento por puesto de la plantilla: 1) el perfil para contratar, 2) lo que tiene que saber,
# 3) su capacitación, 4) lo que prometemos y pedimos, y la constancia de lectura.
# Sueldos y plazas salen de secciones/12_costos.js (una sola fuente). Escribe ../perfiles/*.docx,
# el ZIP con todos y secciones/23_perfiles_datos.js para la página de la app.
import ast, datetime, json, os, re, sys, unicodedata, zipfile
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.join(AQUI, '..')
sys.path.insert(0, AQUI)
from textos_perfiles import PERFILES, COMUN

VERSION, FECHA, FECHA_TXT = 'v1', '2026-09-29', '29 de septiembre de 2026'
NAVY, VERDE, AZUL, GRIS, AMBAR = RGBColor(0x00, 0x12, 0x40), RGBColor(0x0B, 0x7A, 0x55), RGBColor(0x00, 0x39, 0xC8), RGBColor(0x55, 0x5F, 0x78), RGBColor(0xA8, 0x6A, 0x00)

# ── plazas y sueldos desde la app ──
js = open(os.path.join(AQUI, 'secciones', '12_costos.js'), encoding='utf-8').read()
def arreglo(nombre):
    i = js.index(f'const {nombre} = [') + len(f'const {nombre} = ')
    j = js.index('];', i) + 1
    return ast.literal_eval(re.sub(r'//[^\n]*', '', js[i:j]))
PUESTOS = arreglo('PUESTOS'); ESTRUCTURA = arreglo('ESTRUCTURA')
PLAZAS = {n: {'grupo': 'Estructura de proyecto', 'cant': c, 'sueldo': s, 'var': v, 'turno': 'día, con guardia', 'anual': True} for n, c, s, v in ESTRUCTURA}
PLAZAS.update({n: {'grupo': g, 'cant': c, 'sueldo': s, 'var': v, 'turno': t, 'anual': False} for g, n, c, s, v, t in PUESTOS})
faltan = [n for n in PLAZAS if n not in PERFILES]; sobran = [n for n in PERFILES if n not in PLAZAS]
assert not faltan and not sobran, (faltan, sobran)

def dinero(x): return '$' + f'{round(x):,}'
def slug(t):
    t = unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', t.lower()).strip('-')[:60]

# ── piezas de Word ──
def campo_pagina(parrafo):
    r = parrafo.add_run()
    for tipo, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tipo:
            f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), tipo); r._r.append(f)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; r._r.append(it)
def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color); tcPr.append(sh)

class Doc:
    def __init__(self, puesto):
        self.d = Document(); d = self.d
        st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.15
        for nivel, tam, col in ((1, 17, NAVY), (2, 13, AZUL), (3, 11.5, NAVY)):
            h = d.styles[f'Heading {nivel}']; h.font.name = 'Calibri'; h.font.size = Pt(tam); h.font.bold = True; h.font.color.rgb = col
            h.paragraph_format.space_before = Pt(14 if nivel == 1 else 10); h.paragraph_format.space_after = Pt(5)
        s = d.sections[0]
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
        pie = s.footer.paragraphs[0]; pie.text = f'Carreteras de México · Perfil de puesto · {puesto} · {VERSION} · '
        for r in pie.runs: r.font.size = Pt(8); r.font.color.rgb = GRIS
        campo_pagina(pie); pie.runs[-1].font.size = Pt(8)
        cab = s.header.paragraphs[0]; cab.text = 'carreteras.capitaltorreon.com'
        cab.runs[0].font.size = Pt(8); cab.runs[0].font.color.rgb = GRIS
        cp = d.core_properties; cp.author = 'Ing. Ricardo López Reyero'; cp.title = f'Perfil de puesto · {puesto}'
        cp.subject = 'Carreteras de México · manual de contratación e ingreso'; cp.keywords = 'RLR · EYE · 181218'; cp.comments = 'RLR'
        cp.created = cp.modified = datetime.datetime(2026, 9, 29, 7, 0)
    def h(self, t, n=1): return self.d.add_heading(t, n)
    def p(self, t='', negrita=None, tam=None, color=None, cursiva=False, estilo=None):
        par = self.d.add_paragraph(style=estilo)
        # **texto** en negrita dentro de la cadena
        for k, trozo in enumerate(re.split(r'\*\*', t)):
            if not trozo: continue
            r = par.add_run(trozo); r.bold = bool(k % 2) or bool(negrita); r.italic = cursiva
            if tam: r.font.size = Pt(tam)
            if color: r.font.color.rgb = color
        return par
    def viñetas(self, items, casilla=False):
        for it in items: self.p(('☐  ' if casilla else '') + it, estilo=None if casilla else 'List Bullet')
    def tabla(self, cab, filas, anchos=None, encabezado='0039C8'):
        t = self.d.add_table(rows=1, cols=len(cab)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, c in enumerate(cab):
            cel = t.rows[0].cells[i]; cel.text = ''; r = cel.paragraphs[0].add_run(c); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            sombrear(cel, encabezado)
        for f in filas:
            celdas = t.add_row().cells
            for i, v in enumerate(f):
                celdas[i].text = ''
                for k, trozo in enumerate(re.split(r'\*\*', str(v))):
                    if trozo:
                        r = celdas[i].paragraphs[0].add_run(trozo); r.font.size = Pt(9.5); r.bold = bool(k % 2)
        if anchos:
            for fila in t.rows:
                for i, w in enumerate(anchos): fila.cells[i].width = Cm(w)
        self.d.add_paragraph()
        return t
    def cita(self, t, color=VERDE):
        par = self.p(t, negrita=True, tam=11.5, color=color); par.paragraph_format.left_indent = Cm(0.6); par.paragraph_format.space_before = Pt(6)
        return par
    def salto(self): self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

ESQUEMAS = {
    'frente': ('Cada mes, con las cuatro metas del frente', [('40 %', 'Metros colados del frente contra la meta del mes'), ('30 %', 'Calidad medida'), ('20 %', 'Continuidad: minutos de paro y colchones'), ('10 %', 'Seguridad: cero accidentes con tiempo perdido')],
               'Candado: si el mes falla la calidad o la seguridad, la parte de metros no se paga. Primero bien, luego rápido.'),
    'calidad': ('Cada mes, y nada por metros', [('60 %', 'Calidad medida'), ('25 %', 'Seguridad'), ('15 %', 'Tablero y registros completos y a tiempo')],
                'Tu bono no depende del avance a propósito: quien puede rechazar un camión no puede ganar más por dejarlo pasar.'),
    'seguridad': ('Cada mes, y nada por metros', [('60 %', 'Seguridad: cero accidentes con tiempo perdido y cero incidentes con terceros'), ('25 %', 'Señalización y controles auditados sin faltantes'), ('15 %', 'Reportes y pláticas completos y a tiempo')],
                  'Quien protege a la gente no puede ganar más por apurar el frente.'),
    'mando': ('Cada trimestre, con las metas del frente y el margen', [('35 %', 'Metros colados contra el programa'), ('30 %', 'Calidad medida'), ('15 %', 'Continuidad y colchones'), ('10 %', 'Seguridad'), ('10 %', 'Margen del proyecto')],
              'El mismo candado: sin calidad y seguridad no se paga la parte de metros.'),
    'direccion': ('Cada año, pagado en tres partes', [('40 %', 'Al entregar el tramo'), ('30 %', 'Al año 2, si el tramo cumple'), ('30 %', 'Al año 5, si el tramo sigue cumpliendo la garantía (suavidad, agrietamiento, drenes)')],
                  'Cobras cuando la carretera demuestra que dura.'),
    'calidad_dir': ('Cada año, y nada por metros', [('60 %', 'Calidad medida del año'), ('25 %', 'Seguridad'), ('15 %', 'Tablero y trazabilidad completos')],
                    'Si tu bono dependiera de los metros, tu puesto dejaría de existir.'),
}

def generar(num, nombre, P, pl):
    D = Doc(nombre); d = D.d
    # ── portada ──
    D.p('CARRETERAS DE MÉXICO · PERFIL DE PUESTO Y MANUAL DE INGRESO', negrita=True, tam=9, color=VERDE)
    t = D.p(nombre, negrita=True, tam=26, color=NAVY); t.paragraph_format.space_before = Pt(18)
    D.p(P['mision'], tam=12.5, color=GRIS)
    D.p()
    var_txt = f"{pl['var']} % · {dinero(pl['sueldo'] * pl['var'] / 100)} {'al año sobre el sueldo anual' if pl['anual'] else 'al mes'}" if not pl['anual'] else f"{pl['var']} % del sueldo anual · {dinero(pl['sueldo'] * 12 * pl['var'] / 100)}"
    D.tabla(['Dato', ''], [
        ['Línea', P['linea']], ['Reporta a', P['reporta']], ['Grupo', pl['grupo']],
        ['Plazas por frente' if not pl['anual'] else 'Plazas', f"{pl['cant']}" + ('' if not pl['anual'] else ' por proyecto, compartida entre frentes')],
        ['Turno', pl['turno']], ['Sueldo mensual bruto de referencia', dinero(pl['sueldo'])], ['Variable por metas', var_txt],
    ], anchos=[5.5, 11.5], encabezado='001240')
    D.p(f'Versión {VERSION[1:]} · {FECHA_TXT} · Torreón, Coahuila, México', tam=9, color=GRIS)
    D.h('Cómo se usa este documento', 3)
    D.viñetas(['**Para contratar:** la parte 1 dice a quién buscamos, qué preguntar y qué prueba hacer.',
               '**Para la persona contratada:** se entrega completo el día de la oferta. Se lee todo antes del primer turno; las partes 2 y 3 son lo que tiene que saber y cómo la vamos a formar.',
               '**Para cerrar el ingreso:** al final del periodo de capacitación se hacen las preguntas de comprobación y se firma la constancia de lectura.'])
    D.salto()
    # ── carta ──
    D.h('Te damos la bienvenida', 1)
    for par in COMUN['carta'](nombre, P['mision']): D.p(par)
    D.p('**Ing. Ricardo López Reyero**'); D.p('Torreón, Coahuila, México', tam=9.5, color=GRIS)
    D.salto()
    # ── parte 1 ──
    D.h('Parte 1 · El perfil: a quién buscamos', 1)
    D.h('Para qué existe este puesto', 2); D.p(P['mision'])
    if P.get('importa'): D.cita(P['importa'])
    D.h('De qué responde', 2); D.viñetas(P['responde'])
    D.h('Qué puede decidir', 2); D.viñetas(P['autoridad'])
    D.h('Experiencia que pedimos', 2); D.viñetas(P['experiencia'])
    D.h('Escolaridad y certificaciones', 2); D.viñetas(P['escolaridad'])
    D.h('El tipo de persona', 2); D.viñetas(P['persona'])
    D.h('Señales de alerta en la entrevista', 2); D.viñetas(P['alertas'])
    D.h('Preguntas de entrevista', 2)
    D.tabla(['Pregunta', 'Lo que buscamos en la respuesta'], P['preguntas'], anchos=[7.5, 9.5])
    D.h('Prueba práctica', 2); D.p(P['prueba'])
    D.h('Sueldo y variable', 2)
    ev, partes, candado = ESQUEMAS[P['esquema']]
    D.p(f"Sueldo mensual bruto de referencia: **{dinero(pl['sueldo'])}**, más prestaciones de ley desde el primer día. " +
        (f"Variable meta: **{pl['var']} % del sueldo anual** ({dinero(pl['sueldo'] * 12 * pl['var'] / 100)})." if pl['anual'] else f"Variable meta: **{pl['var']} % del sueldo** ({dinero(pl['sueldo'] * pl['var'] / 100)} al mes); con la meta cumplida el mes suma **{dinero(pl['sueldo'] * (1 + pl['var'] / 100))}**."))
    D.p(f'**Cómo se paga:** {ev.lower()}.')
    D.tabla(['Peso', 'Meta'], partes, anchos=[2.5, 14.5])
    D.p('**Tus indicadores:**'); D.viñetas(P['indicadores'])
    D.cita(candado, AMBAR)
    D.salto()
    # ── parte 2 ──
    D.h('Parte 2 · Lo que tienes que saber', 1)
    D.h('Un día normal en tu puesto', 2)
    D.tabla(['Cuándo', 'Qué pasa'], P['dia'], anchos=[3.5, 13.5])
    D.h('Lo específico de tu puesto', 2); D.viñetas(P['saber'])
    D.h('Los números que tienes que traer en la cabeza', 2)
    D.tabla(['Número', 'Qué significa'], P['numeros'], anchos=[4, 13])
    D.h('Los errores que cuestan 50 años', 2); D.viñetas(P['errores'])
    D.h('Lo que todos saben en esta obra', 2)
    COMUN['todos'](D)
    D.salto()
    # ── parte 3 ──
    D.h('Parte 3 · Tu capacitación', 1)
    D.p('Nadie empieza solo ni a velocidad plena. Así es tu entrada, con quién y cómo sabemos que ya estás listo para trabajar por tu cuenta:')
    D.tabla(['Cuándo', 'Qué haces y con quién'], [COMUN['dia1']] + P['capacitacion'], anchos=[3.8, 13.2])
    D.h('Certificaciones y formación', 2); D.viñetas(P['certificaciones'])
    D.h('Preguntas de comprobación', 2)
    D.p('Al terminar tu periodo de ingreso, tu jefe directo te hace estas preguntas. Tienes que poder contestarlas sin leer.')
    for k, q in enumerate(COMUN['comprobacion'] + P['comprobacion'], 1): D.p(f'☐  {k}. {q}')
    D.h('Tu siguiente escalón', 2); D.p(P['siguiente'])
    D.salto()
    # ── parte 4 ──
    D.h('Parte 4 · Lo que te prometemos y lo que te pedimos', 1)
    D.h('Lo que te prometemos', 2); D.viñetas(COMUN['prometemos'])
    D.h('Lo que te pedimos', 2); D.viñetas(COMUN['pedimos'] + P.get('pedimos', []))
    D.h('Constancia de lectura', 2)
    D.p(f'Leí completo este documento del puesto **{nombre}**, lo entendí y resolví mis dudas con mi jefe directo. Sé que puedo detener el trabajo por seguridad sin consecuencias, que en esta obra nunca se añade agua al concreto y que las desviaciones sólo existen si están firmadas.')
    D.p(); D.p()
    D.tabla(['', 'Persona contratada', f"Jefe directo ({P['reporta']})"], [['Nombre', '', ''], ['Firma', '', ''], ['Fecha', '', '']], anchos=[3, 7, 7], encabezado='001240')
    return d

def main():
    carpeta = os.path.join(RAIZ, 'perfiles'); os.makedirs(carpeta, exist_ok=True)
    for f in os.listdir(carpeta):
        if f.endswith('.docx') or f.endswith('.zip'): os.remove(os.path.join(carpeta, f))
    datos, zipnombre = [], 'perfiles-de-puesto.zip'
    orden = [n for n, *_ in ESTRUCTURA] + [p[1] for p in PUESTOS]
    with zipfile.ZipFile(os.path.join(carpeta, zipnombre), 'w', zipfile.ZIP_DEFLATED) as z:
        for num, nombre in enumerate(orden, 1):
            P, pl = PERFILES[nombre], PLAZAS[nombre]
            archivo = f'{num:02d}-{slug(nombre)}.docx'
            descarga = f'Perfil de puesto · {nombre} · {VERSION} · {FECHA}.docx'
            ruta = os.path.join(carpeta, archivo)
            generar(num, nombre, P, pl).save(ruta)
            z.write(ruta, f'Perfiles de puesto · Carreteras de México/{num:02d} · {nombre} · {VERSION}.docx')
            datos.append({'n': nombre, 'g': pl['grupo'], 'c': pl['cant'], 't': pl['turno'], 's': pl['sueldo'], 'v': pl['var'], 'an': pl['anual'],
                          'r': P['reporta'], 'l': P['linea'], 'm': P['mision'], 'x': P['experiencia'][0], 'a': 'perfiles/' + archivo, 'd': descarga,
                          'kb': round(os.path.getsize(ruta) / 1024)})
    kbzip = round(os.path.getsize(os.path.join(carpeta, zipnombre)) / 1024)
    js_out = ('/* ── Perfiles de puesto: índice generado por herramientas/perfiles.py (no editar a mano) — RLR ── */\n'
              f'const PERFILES_DOC = {json.dumps(datos, ensure_ascii=False)};\n'
              f'const PERFILES_ZIP = {json.dumps({"a": "perfiles/" + zipnombre, "d": f"Perfiles de puesto · Carreteras de México · {VERSION} · {FECHA}.zip", "kb": kbzip}, ensure_ascii=False)};\n')
    open(os.path.join(AQUI, 'secciones', '23_perfiles_datos.js'), 'w', encoding='utf-8').write(js_out)
    print(len(datos), 'perfiles ·', sum(x['kb'] for x in datos), 'KB · ZIP', kbzip, 'KB')

if __name__ == '__main__':
    main()
